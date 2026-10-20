import math
import re
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OLIVE_PRIMARY = RGBColor(0x4B, 0x5D, 0x2E)
OLIVE_SECONDARY = RGBColor(0x8A, 0x9A, 0x5B)
OLIVE_ACCENT = RGBColor(0xD4, 0xC5, 0x8E)
OLIVE_LIGHT = RGBColor(0xF4, 0xF2, 0xE6)
TEXT_DARK = RGBColor(0x2B, 0x2B, 0x2B)
TEXT_MUTED = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN_OK = RGBColor(0x3C, 0x7A, 0x3E)
RED_KO = RGBColor(0xB3, 0x40, 0x2F)
AMBER = RGBColor(0xC7, 0x8A, 0x1A)

OUTPUT_DIR = Path(__file__).resolve().parents[2] / "generated_decks"
OUTPUT_DIR.mkdir(exist_ok=True)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

PAGE_CAPACITY = 13.0      # "lignes" disponibles par diapositive de contenu
CHARS_PER_LINE = 95
MATRIX_ROWS_PER_SLIDE = 6


# ---------------------------------------------------------------- Markdown -> texte propre
def _strip_markup(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    return re.sub(r"[*`]+", "", text).strip()


def _parse_runs(text: str) -> list[tuple[str, bool]]:
    """Convertit **gras** en runs réellement gras; supprime tout autre symbole Markdown résiduel."""
    runs = []
    for part in re.split(r"(\*\*[^*]+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True))
        else:
            cleaned = re.sub(r"[*`]+", "", part).replace("__", "")
            if cleaned:
                runs.append((cleaned, False))
    return runs


def _split_long(text: str, max_chars: int = 600) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks, current = [], ""
    for s in sentences:
        if current and len(current) + len(s) > max_chars:
            chunks.append(current.strip())
            current = ""
        current += s + " "
    if current.strip():
        chunks.append(current.strip())
    return chunks


def _parse_section(markdown: str, fallback_title: str) -> tuple[str, list[dict]]:
    title = fallback_title
    blocks: list[dict] = []
    title_found = False

    for line in markdown.splitlines():
        s = line.strip()
        if not s or re.fullmatch(r"[-*_]{3,}", s):
            continue
        if s.startswith("#"):
            heading = _strip_markup(s.lstrip("#").strip())
            heading = re.sub(r"^\d+[.)]\s*", "", heading)
            if not title_found and heading:
                title, title_found = heading, True
            continue
        if s.startswith("|"):
            if re.fullmatch(r"[|\-\s:]+", s):
                continue
            cells = [c.strip() for c in s.strip("|").split("|") if c.strip()]
            blocks.append({"type": "para", "text": " — ".join(cells)})
            continue

        m = re.match(r"^([-*•]|\d+[.)])\s+(.*)", s)
        if m:
            for piece in _split_long(m.group(2)):
                blocks.append({"type": "bullet", "text": piece})
        else:
            for piece in _split_long(s):
                blocks.append({"type": "para", "text": piece})
    return title, blocks


def _paginate(blocks: list[dict]) -> list[list[dict]]:
    pages, current, used = [], [], 0.0
    for b in blocks:
        weight = math.ceil(len(_strip_markup(b["text"])) / CHARS_PER_LINE) + 0.4
        if current and used + weight > PAGE_CAPACITY:
            pages.append(current)
            current, used = [], 0.0
        current.append(b)
        used += weight
    if current:
        pages.append(current)
    return pages


def _display(value, default: str = "À confirmer") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return default if text.lower() in ("", "unknown", "none", "null") else text


# ---------------------------------------------------------------- Construction du deck
class _Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.count = 0

    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.count += 1
        return slide

    @staticmethod
    def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    @staticmethod
    def _add_runs(paragraph, text, size, color, bold=False):
        for chunk, is_bold in _parse_runs(text):
            run = paragraph.add_run()
            run.text = chunk
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold or is_bold
            run.font.color.rgb = color

    def _text(self, slide, x, y, w, h, text, size=16, color=TEXT_DARK, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        self._add_runs(p, text, size, color, bold)
        return box

    def _footer(self, slide):
        self._text(slide, Inches(0.8), Inches(7.0), Inches(8), Inches(0.35),
                   "OliveSoft  •  Proposition confidentielle", size=10, color=TEXT_MUTED)
        self._text(slide, Inches(11.8), Inches(7.0), Inches(1.0), Inches(0.35),
                   str(self.count), size=10, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)

    def _title_bar(self, slide, title):
        self._rect(slide, 0, 0, Inches(0.3), SLIDE_H, OLIVE_PRIMARY)
        self._text(slide, Inches(0.8), Inches(0.35), Inches(11.8), Inches(1.0), title,
                   size=28, color=OLIVE_PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        self._rect(slide, Inches(0.8), Inches(1.4), Inches(1.6), Inches(0.06), OLIVE_ACCENT)

    # ---- slides
    def title_slide(self, tender: dict):
        s = self._blank()
        self._rect(s, 0, 0, SLIDE_W, SLIDE_H, OLIVE_PRIMARY)
        self._rect(s, Inches(0.8), Inches(2.2), Inches(1.2), Inches(0.08), OLIVE_ACCENT)
        self._text(s, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.2),
                   "Proposition commerciale", size=46, color=WHITE, bold=True)
        self._text(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(1.3),
                   _display(tender.get("title"), "Appel d'offres"), size=22, color=OLIVE_ACCENT)
        org = _display(tender.get("issuing_organization"), "")
        if org:
            self._text(s, Inches(0.8), Inches(5.2), Inches(11.5), Inches(0.5),
                       f"Préparée pour {org}", size=18, color=WHITE)
        self._text(s, Inches(0.8), Inches(5.8), Inches(11.5), Inches(0.5),
                   f"OliveSoft  •  {datetime.now().strftime('%d/%m/%Y')}", size=14, color=OLIVE_ACCENT)

    def agenda_slide(self, titles: list[str]):
        s = self._blank()
        self._title_bar(s, "Sommaire")
        box = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, t in enumerate(titles):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            self._add_runs(p, f"{i + 1:02d}   ", 20, OLIVE_SECONDARY, bold=True)
            self._add_runs(p, t, 20, TEXT_DARK)
        self._footer(s)

    def summary_slide(self, tender: dict, matrix: list[dict]):
        s = self._blank()
        self._title_bar(s, "L'opportunité en un coup d'œil")

        total = len(matrix)
        covered = sum(1 for m in matrix if m.get("covered"))
        coverage = f"{round(100 * covered / total)} %" if total else "À confirmer"
        qual = tender.get("qualification_score")
        cards = [
            ("Adéquation globale", f"{round(qual)}/100" if qual is not None else "À confirmer"),
            ("Exigences couvertes", f"{coverage}" + (f" ({covered}/{total})" if total else "")),
            ("Budget", _display(tender.get("budget_display") or tender.get("estimated_budget"))),
            ("Date limite", _display(tender.get("deadline"))),
        ]
        card_w, gap, x0 = Inches(2.95), Inches(0.15), Inches(0.65)
        for i, (label, value) in enumerate(cards):
            x = x0 + i * (card_w + gap)
            self._rect(s, x, Inches(1.8), card_w, Inches(1.7), OLIVE_LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
            self._rect(s, x, Inches(1.8), Inches(0.1), Inches(1.7), OLIVE_SECONDARY)
            self._text(s, x + Inches(0.25), Inches(1.9), card_w - Inches(0.35), Inches(0.5),
                       label, size=13, color=TEXT_MUTED)
            self._text(s, x + Inches(0.25), Inches(2.4), card_w - Inches(0.35), Inches(1.0),
                       value, size=22, color=OLIVE_PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

        box = s.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.8), Inches(2.6))
        tf = box.text_frame
        tf.word_wrap = True
        rows = [
            ("Organisation", _display(tender.get("issuing_organization"))),
            ("Secteur", _display(tender.get("sector"))),
            ("Objet", _display(tender.get("summary"), "")[:380]),
        ]
        first = True
        for label, value in rows:
            if not value:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            self._add_runs(p, f"**{label} :** {value}", 15, TEXT_DARK)

        if tender.get("urgent_deadline_alert"):
            self._rect(s, Inches(0.8), Inches(6.35), Inches(11.8), Inches(0.5), RED_KO, MSO_SHAPE.ROUNDED_RECTANGLE)
            self._text(s, Inches(1.0), Inches(6.35), Inches(11.4), Inches(0.5),
                       "Échéance proche : réponse à prioriser", size=15, color=WHITE, bold=True,
                       anchor=MSO_ANCHOR.MIDDLE)
        self._footer(s)

    def content_slide(self, title: str, blocks: list[dict], continued: bool):
        s = self._blank()
        self._title_bar(s, title + (" (suite)" if continued else ""))
        box = s.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.8), Inches(5.0))
        tf = box.text_frame
        tf.word_wrap = True
        for i, b in enumerate(blocks):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(9)
            if b["type"] == "bullet":
                self._add_runs(p, "•  ", 16, OLIVE_SECONDARY, bold=True)
            self._add_runs(p, b["text"], 16, TEXT_DARK)
        self._footer(s)

    def _set_cell(self, cell, text, size=12, color=TEXT_DARK, bold=False, fill=None, align=PP_ALIGN.LEFT):
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = align
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        self._add_runs(p, text, size, color, bold)
        if fill is not None:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill

    def matrix_slides(self, matrix: list[dict]):
        pages = [matrix[i:i + MATRIX_ROWS_PER_SLIDE] for i in range(0, len(matrix), MATRIX_ROWS_PER_SLIDE)]
        for idx, rows in enumerate(pages):
            s = self._blank()
            self._title_bar(s, "Matrice de couverture des exigences" + (" (suite)" if idx else ""))
            shape = s.shapes.add_table(len(rows) + 1, 4, Inches(0.8), Inches(1.7), Inches(11.8), Inches(4.9))
            table = shape.table
            for c, w in enumerate([4.6, 1.5, 1.0, 4.7]):
                table.columns[c].width = Inches(w)

            for c, header in enumerate(["Exigence", "Couverture", "Score", "Preuve / référence OliveSoft"]):
                self._set_cell(table.cell(0, c), header, 13, WHITE, True, OLIVE_PRIMARY,
                               PP_ALIGN.CENTER if c in (1, 2) else PP_ALIGN.LEFT)

            for r, item in enumerate(rows, start=1):
                band = OLIVE_LIGHT if r % 2 == 0 else WHITE
                covered = bool(item.get("covered"))
                evidence = _display(item.get("concrete_evidence"), "")
                ref = _display(item.get("best_document_title"), "")
                proof = evidence or ref or "À approfondir avec le client"
                self._set_cell(table.cell(r, 0), _strip_markup(item["requirement_text"])[:130], 12, fill=band)
                self._set_cell(table.cell(r, 1), "Couverte" if covered else "À approfondir", 12, WHITE, True,
                               GREEN_OK if covered else AMBER, PP_ALIGN.CENTER)
                self._set_cell(table.cell(r, 2), f"{round(item['score'])}/100", 12, fill=band, align=PP_ALIGN.CENTER)
                self._set_cell(table.cell(r, 3), _strip_markup(proof)[:150], 11, fill=band)
            self._footer(s)

    def closing_slide(self):
        s = self._blank()
        self._rect(s, 0, 0, SLIDE_W, SLIDE_H, OLIVE_PRIMARY)
        self._rect(s, Inches(0.8), Inches(2.9), Inches(1.2), Inches(0.08), OLIVE_ACCENT)
        self._text(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(1.2),
                   "Merci pour votre confiance", size=42, color=WHITE, bold=True)
        self._text(s, Inches(0.8), Inches(4.4), Inches(11.5), Inches(0.8),
                   "Nous restons à votre disposition pour échanger sur cette proposition.",
                   size=18, color=OLIVE_ACCENT)


def generate_pptx(tender: dict, proposal_sections: dict, matrix: list[dict], filename: str) -> Path:
    deck = _Deck()

    parsed = []
    for key, content in proposal_sections.items():
        fallback = key.replace("_", " ").capitalize()
        title, blocks = _parse_section(content or "", fallback)
        if blocks:
            parsed.append((title, blocks))

    deck.title_slide(tender)
    deck.summary_slide(tender, matrix)
    deck.agenda_slide([t for t, _ in parsed] + (["Matrice de couverture des exigences"] if matrix else []))

    for title, blocks in parsed:
        for page_idx, page in enumerate(_paginate(blocks)):
            deck.content_slide(title, page, continued=page_idx > 0)

    if matrix:
        deck.matrix_slides(matrix)

    deck.closing_slide()

    output_path = OUTPUT_DIR / f"{filename}.pptx"
    deck.prs.save(output_path)
    return output_path